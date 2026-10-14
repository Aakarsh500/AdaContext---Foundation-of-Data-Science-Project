"""Build the submission-ready AdaContext FDS project outline as a DOCX."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "Adaptive_Context_Selection_FDS_Project_Outline.docx"

BLUE = RGBColor(46, 116, 181)
NAVY = RGBColor(31, 77, 120)
INK = RGBColor(33, 41, 52)
GRAY = RGBColor(91, 99, 112)
MUTED = RGBColor(118, 126, 138)
WHITE = RGBColor(255, 255, 255)
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F4F6F9"
VERY_LIGHT = "FAFBFC"
GRID = "D7DBE2"

PAGE_WIDTH_DXA = 12240
MARGIN_DXA = 1440
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, *, size=None, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_page(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def configure_styles(doc: Document):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.widow_control = True

    heading_specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, NAVY, 8, 4),
    }
    for style_name, (size, color, before, after) in heading_specs.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.font.color.rgb = INK
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.208


def add_numbering_instance(doc: Document, abstract_id: int) -> int:
    numbering = doc.part.numbering_part.element
    num_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    num_id = max(num_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return num_id


def create_numbering(doc: Document):
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))
    ]

    def add_definition(fmt: str, text: str, font: str = "Calibri") -> int:
        abstract_id = max(abstract_ids, default=-1) + 1
        abstract_ids.append(abstract_id)
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)

        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        level.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        level.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        level.append(lvl_text)
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        level.append(lvl_jc)
        ppr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ppr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "279")
        ppr.append(ind)
        level.append(ppr)
        rpr = OxmlElement("w:rPr")
        rfonts = OxmlElement("w:rFonts")
        rfonts.set(qn("w:ascii"), font)
        rfonts.set(qn("w:hAnsi"), font)
        rpr.append(rfonts)
        level.append(rpr)
        abstract.append(level)
        # OOXML requires every abstractNum before the first concrete num instance.
        first_num = numbering.find(qn("w:num"))
        if first_num is None:
            numbering.append(abstract)
        else:
            numbering.insert(list(numbering).index(first_num), abstract)
        return abstract_id

    bullet_abstract_id = add_definition("bullet", "•")
    decimal_abstract_id = add_definition("decimal", "%1.")
    bullet_num_id = add_numbering_instance(doc, bullet_abstract_id)
    return bullet_num_id, decimal_abstract_id


def apply_numbering(paragraph, num_id: int):
    ppr = paragraph._p.get_or_add_pPr()
    existing = ppr.find(qn("w:numPr"))
    if existing is not None:
        ppr.remove(existing)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_pr.append(ilvl)
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(num)
    ppr.append(num_pr)


def add_list_item(doc, text: str, num_id: int, *, bold_lead: str | None = None):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    apply_numbering(p, num_id)
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.194)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_lead) :])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_numbered_item(doc, text: str, num_id: int, *, bold_lead: str | None = None):
    p = doc.add_paragraph(style="List Number")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    apply_numbering(p, num_id)
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.194)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_lead) :])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_body(doc, text: str, *, bold_lead: str | None = None, italic=False, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_lead) :])
        set_run_font(r, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic)
    return p


def set_paragraph_box(paragraph, *, fill=LIGHT_BLUE, border=NAVY):
    border_value = str(border)
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), border_value)
    borders.append(left)
    ppr.append(borders)


def add_callout(doc, label: str, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.2
    set_paragraph_box(p)
    r = p.add_run(label + "  ")
    set_run_font(r, bold=True, color=NAVY, size=10.5)
    r = p.add_run(text)
    set_run_font(r, color=INK, size=10.5)
    return p


def set_cell_margins(cell, *, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def configure_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), GRID)
        borders.append(node)


def add_table(
    doc,
    headers,
    rows,
    widths_dxa,
    *,
    font_size=9.2,
    first_col_bold=False,
    keep_together=True,
):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.style = "Table Grid"

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    configure_table_borders(table)

    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, (cell, value) in enumerate(zip(header.cells, headers)):
        set_cell_width(cell, widths_dxa[i])
        set_cell_margins(cell)
        shade_cell(cell, LIGHT_GRAY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(str(value))
        set_run_font(r, size=font_size, bold=True, color=NAVY)

    for row_idx, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, (cell, value) in enumerate(zip(row.cells, values)):
            set_cell_width(cell, widths_dxa[i])
            set_cell_margins(cell)
            if row_idx % 2 == 1:
                shade_cell(cell, VERY_LIGHT)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, bold=(first_col_bold and i == 0), color=INK)

    # Keep compact tables together instead of leaving a header and one row stranded.
    if keep_together:
        all_rows = table.rows
        for row in all_rows[:-1]:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = True

    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(4)
    after.paragraph_format.space_after = Pt(2)
    return table


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_run_font(run, size=8.5, color=MUTED)


def configure_running_header(header, *, show_text: bool):
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    if not show_text:
        return
    r = p.add_run("ADACONTEXT")
    set_run_font(r, size=8.5, bold=True, color=BLUE)
    r = p.add_run("   |   FOUNDATION OF DATA SCIENCE")
    set_run_font(r, size=8.5, color=MUTED)


def configure_running_footer(footer, *, full: bool):
    footer.is_linked_to_previous = False
    for existing in list(footer.tables):
        footer._element.remove(existing._tbl)
    p = footer.paragraphs[0]
    p.text = ""
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    if not full:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("Page ")
        set_run_font(r, size=8.5, color=MUTED)
        add_page_number(p)
        return
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.style = "Table Grid"
    widths = [7020, 2340]
    for i, cell in enumerate(table.rows[0].cells):
        set_cell_width(cell, widths[i])
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_p = cell.paragraphs[0]
        cell_p.paragraph_format.space_after = Pt(0)
        cell_p.paragraph_format.line_spacing = 1
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    left_p = table.cell(0, 0).paragraphs[0]
    left_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = left_p.add_run("Adaptive Context Selection — Project Outline")
    set_run_font(r, size=8.5, color=MUTED)
    right_p = table.cell(0, 1).paragraphs[0]
    right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = right_p.add_run("Page ")
    set_run_font(r, size=8.5, color=MUTED)
    add_page_number(right_p)


def configure_header_footer(section):
    section.different_first_page_header_footer = True
    configure_running_header(section.header, show_text=False)
    configure_running_header(section.even_page_header, show_text=False)
    configure_running_header(section.first_page_header, show_text=False)
    configure_running_footer(section.footer, full=False)
    configure_running_footer(section.even_page_footer, full=False)
    configure_running_footer(section.first_page_footer, full=False)


def add_heading(doc, text: str, level: int):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    return p


def add_cover(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(48)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("FOUNDATION OF DATA SCIENCE  •  ASSIGNMENT 1")
    set_run_font(r, size=10, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Adaptive Context Selection")
    set_run_font(r, size=30, bold=True, color=NAVY)
    r.add_break()
    r2 = p.add_run("for Long-Context LLMs")
    set_run_font(r2, size=30, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run(
        "AdaContext: Budget-Aware Hybrid Evidence Selection for\n"
        "Long-Document Question Answering"
    )
    set_run_font(r, size=14, color=BLUE)

    rule = doc.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rule.paragraph_format.left_indent = Inches(1.35)
    rule.paragraph_format.right_indent = Inches(1.35)
    rule.paragraph_format.space_after = Pt(22)
    ppr = rule._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2E74B5")
    borders.append(bottom)
    ppr.append(borders)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("PROJECT OUTLINE AND TECHNICAL DOCUMENTATION")
    set_run_font(r, size=11, bold=True, color=INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(34)
    r = p.add_run("Completed CPU-Compatible Implementation  •  August 2026")
    set_run_font(r, size=10.5, color=GRAY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Inches(0.55)
    p.paragraph_format.right_indent = Inches(0.55)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = 1.2
    set_paragraph_box(p, fill=LIGHT_BLUE, border="2E74B5")
    r = p.add_run("Core research question\n")
    set_run_font(r, size=10, bold=True, color=BLUE)
    r = p.add_run(
        "Can a learned, question-specific context budget preserve answer quality while "
        "substantially reducing the tokens and time required for long-document question answering?"
    )
    set_run_font(r, size=12, italic=True, color=INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("Python  •  Streamlit  •  Scikit-learn  •  Sentence Transformers")
    set_run_font(r, size=9.5, color=MUTED)

    doc.add_page_break()


def build_document():
    doc = Document()
    configure_page(doc.sections[0])
    configure_styles(doc)
    bullet_id, decimal_abstract_id = create_numbering(doc)
    doc.settings.odd_and_even_pages_header_footer = True
    configure_header_footer(doc.sections[0])

    props = doc.core_properties
    props.title = "Adaptive Context Selection for Long-Context LLMs"
    props.subject = "Foundation of Data Science project outline and technical documentation"
    props.author = "AdaContext FDS Project"
    props.keywords = "data science, long context, retrieval, adaptive context, LLM, QASPER"
    props.comments = "Generated from the completed AdaContext implementation."

    add_cover(doc)

    add_heading(doc, "Executive Abstract", 1)
    add_body(
        doc,
        "Modern large language models can accept very long inputs, but passing an entire document "
        "for every question is expensive, slow, and vulnerable to evidence dilution and the "
        "lost-in-the-middle effect. AdaContext treats context construction as a data-science "
        "allocation problem: retrieve candidate evidence, estimate question difficulty, assign a "
        "question-specific context budget, and select the smallest useful set of diverse chunks."
    )
    add_body(
        doc,
        "The completed system combines BM25 lexical retrieval, semantic retrieval, reciprocal-rank "
        "fusion, optional cross-encoder reranking, feature engineering, a learned budget classifier, "
        "and a token-aware selector. It includes reproducible preprocessing, EDA, six evaluation "
        "strategies, bootstrap uncertainty estimates, a lost-in-the-middle stress test, monitoring "
        "logs, a command-line interface, and a Streamlit application. It runs on a normal CPU and "
        "also supports custom questions over bundled, QASPER, or uploaded documents."
    )
    add_callout(
        doc,
        "Project contribution",
        "The novelty is not a new foundation model. It is a measurable policy that learns how much "
        "context a question needs and exposes the quality–cost trade-off with auditable evidence."
    )

    add_heading(doc, "Project at a Glance", 2)
    for text, lead in [
        ("Problem: fixed-size or full-document prompting wastes tokens and can bury relevant evidence.", "Problem:"),
        ("Method: hybrid retrieval plus a learned three-class context-budget controller and diversity-aware selection.", "Method:"),
        ("Evaluation: answer quality, evidence coverage, context precision, token cost, latency, confidence intervals, and position stress tests.", "Evaluation:"),
        ("Deliverable: a complete Python package, datasets, notebooks, reports, tests, monitoring logs, and interactive Streamlit demo.", "Deliverable:"),
    ]:
        add_list_item(doc, text, bullet_id, bold_lead=lead)

    add_heading(doc, "1. Problem Definition", 1)
    add_heading(doc, "1.1 Background and Motivation", 2)
    add_body(
        doc,
        "Long-context LLMs create the impression that every available page should be included in the "
        "prompt. In practice, context length has direct computational cost, long prompts increase "
        "response time, and irrelevant passages can compete with the evidence needed for the answer. "
        "A fixed top-k retriever is better than full-context prompting, but it still assumes that a "
        "simple factual question and a broad multi-evidence question need the same budget."
    )
    add_body(
        doc,
        "AdaContext instead predicts the required context size separately for each question. The "
        "independent variable is the selection strategy; the main dependent variables are answer "
        "quality, evidence recall, token use, and latency. Document length, question type, generator, "
        "and retrieval backend are recorded as controls or grouping variables."
    )

    add_heading(doc, "1.2 Research Objective", 2)
    add_callout(
        doc,
        "Optimization view",
        "For question q, all chunks C, and selected subset S, the project seeks S* that maximizes "
        "Q(q,S) − λ[T(S)/T(C)] − μL(S), where Q is answer quality, T is token cost, and L is latency."
    )
    add_body(doc, "The operational goals are to:")
    objectives = [
        "retrieve evidence using complementary lexical and semantic signals;",
        "predict a simple, moderate, or complex context budget without using the correct answer;",
        "select relevant yet non-redundant chunks within that budget;",
        "expand the budget automatically when retrieval confidence is low; and",
        "match or improve answer quality while reducing context tokens and processing time."
    ]
    for item in objectives:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "1.3 Research Hypotheses", 2)
    hypotheses = [
        "H1: Hybrid retrieval achieves higher gold-evidence recall than BM25 or dense retrieval alone.",
        "H2: The adaptive policy uses fewer context tokens than fixed hybrid retrieval at comparable answer quality.",
        "H3: Confidence-based expansion reduces failures for ambiguous, broad, or multi-evidence questions.",
        "H4: Adaptive evidence selection is less sensitive to evidence position and document length than naive long-context prompting."
    ]
    for item in hypotheses:
        add_list_item(doc, item, bullet_id, bold_lead=item[:3])

    add_heading(doc, "1.4 Scope", 2)
    add_body(
        doc,
        "The project focuses on question answering over a known collection of documents. It does not "
        "train a large language model. The learned component is a lightweight classifier, which keeps "
        "the work reproducible on a laptop while still demonstrating feature engineering, supervised "
        "learning, validation, statistical comparison, deployment, and monitoring."
    )

    add_heading(doc, "2. Data Collection", 1)
    add_heading(doc, "2.1 Data Sources", 2)
    add_body(
        doc,
        "Three complementary data sources support development, research evaluation, and live use. "
        "Gold evidence is important because it allows retrieval quality to be measured separately from "
        "the quality of the final answer generator."
    )
    add_table(
        doc,
        ["Source", "Purpose", "Scale / Format"],
        [
            ["QASPER", "Primary research dataset with expert answers and supporting evidence.", "5,049 questions across 1,585 NLP papers; official train, validation, and test Parquet files."],
            ["Bundled benchmark", "Offline software verification and classroom demonstration.", "4 fictional structured reports, 24 manually authored questions, answers, evidence, difficulty, and splits."],
            ["User documents", "Custom-question workflow in the application.", "PDF, TXT, Markdown, and HTML uploads; any question grounded in the indexed content."],
        ],
        [1700, 3860, 3800],
        first_col_bold=True,
    )

    add_heading(doc, "2.2 Question and Answer Records", 2)
    add_body(doc, "Each evaluation record stores:")
    for item in [
        "a document identifier and split;",
        "the natural-language question and one or more accepted answers;",
        "answer type and available difficulty metadata, plus gold evidence text or evidence locations; and",
        "derived mappings from evidence spans to chunk identifiers."
    ]:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "2.3 Split and Leakage Controls", 2)
    add_body(
        doc,
        "QASPER uses document-level train, validation, and test splits, preventing questions from the "
        "same paper from leaking across partitions. The controller is trained only on permitted "
        "training records; configuration is selected on validation data; and held-out test results are "
        "reported once. Gold evidence creates labels and evaluation metrics but is never available to "
        "the online controller."
    )

    add_heading(doc, "3. Data Preprocessing", 1)
    add_heading(doc, "3.1 Preprocessing Pipeline", 2)
    steps = [
        ("Extract text. ", "Read PDF pages or structured text while retaining document and page identifiers."),
        ("Normalize content. ", "Apply Unicode normalization, remove duplicated whitespace, dehyphenate PDF line breaks, and clean lightweight markup."),
        ("Preserve structure. ", "Retain titles, section headings, page numbers, source paths, and original chunk order."),
        ("Create chunks. ", "Split on section and sentence boundaries where possible, targeting 220 approximate tokens with 40-token overlap; enforce 45–330-token bounds."),
        ("Align gold evidence. ", "Map annotated evidence to chunks using normalized containment and token-level similarity."),
        ("Build reusable indexes. ", "Compute BM25 statistics and dense or TF-IDF/SVD representations once per corpus."),
    ]
    number_id = add_numbering_instance(doc, decimal_abstract_id)
    for lead, rest in steps:
        add_numbered_item(doc, lead + rest, number_id, bold_lead=lead)

    add_heading(doc, "3.2 Chunk Metadata", 2)
    add_body(
        doc,
        "Every chunk carries its source document, section, page, order, cleaned text, and approximate "
        "token count. This metadata supports citations, neighbour recovery, section-diversity features, "
        "audit logs, and reconstruction of document order."
    )

    add_heading(doc, "3.3 Data-Quality Risks", 2)
    for text, lead in [
        ("PDF extraction: multi-column pages, headers, or broken words can distort sentence boundaries.", "PDF extraction:"),
        ("Evidence boundaries: an annotation may span two chunks, so overlap and approximate matching are required.", "Evidence boundaries:"),
        ("Scientific-paper noise: references, formulas, and bibliography tokens can affect lexical scores.", "Scientific-paper noise:"),
        ("Unanswerable questions: the system must avoid fabricating an answer when retrieved evidence is insufficient.", "Unanswerable questions:"),
    ]:
        add_list_item(doc, text, bullet_id, bold_lead=lead)

    add_heading(doc, "4. Exploratory Data Analysis (EDA)", 1)
    add_heading(doc, "4.1 Questions Addressed by EDA", 2)
    add_body(
        doc,
        "EDA checks whether the corpus, chunking scheme, and retrieval signals are suitable before "
        "training the budget controller. It also tests the central assumption that question complexity "
        "and retrieval uncertainty provide useful information about the amount of context required."
    )
    for item in [
        "How many chunks and tokens does each document produce?",
        "How are questions distributed by difficulty and answer type?",
        "How many gold-evidence chunks are needed per question?",
        "Do gold chunks score higher than non-gold chunks?",
        "When do BM25 and dense retrieval agree or disagree?",
        "Do entropy, score gaps, and section dispersion increase for multi-evidence questions?"
    ]:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "4.2 Automated EDA Outputs", 2)
    add_body(
        doc,
        "The EDA script writes audit-ready CSV files plus interactive Plotly charts for chunks per "
        "document, chunk-token distribution, gold-versus-non-gold scores, and uncertainty versus "
        "evidence count. Raw tables remain available for manual checking rather than relying only on "
        "visual summaries."
    )

    add_heading(doc, "4.3 Bundled Benchmark Profile", 2)
    add_table(
        doc,
        ["Metric", "Observed value"],
        [
            ["Documents", "4"],
            ["Questions", "24"],
            ["Chunks", "20"],
            ["Mean chunks per document", "5.0"],
            ["Median chunk size", "98.5 approximate tokens"],
            ["Mean gold-evidence chunks per question", "1.125"],
            ["Offline embedding backend", "TF-IDF/SVD"],
        ],
        [5600, 3760],
        first_col_bold=True,
    )
    add_callout(
        doc,
        "Interpretation caution",
        "The bundled data validates the software and presentation flow. It is too small for scientific "
        "claims; final conclusions should be based on the held-out QASPER experiment."
    )

    add_heading(doc, "5. Model Building", 1)
    add_heading(doc, "5.1 End-to-End Architecture", 2)
    add_body(
        doc,
        "The system is a multi-stage evidence-allocation pipeline. Candidate retrieval is separated "
        "from budget prediction so that the project can test whether adaptation adds value beyond a "
        "strong fixed hybrid retriever."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_together = True
    r = p.add_run(
        "DOCUMENTS  →  STRUCTURED CHUNKS  →  HYBRID RETRIEVAL  →  RERANKING\n"
        "QUESTION  →  FEATURES  →  BUDGET CONTROLLER  →  ADAPTIVE SELECTOR  →  GROUNDED ANSWER + TRACE"
    )
    set_run_font(r, size=10, bold=True, color=NAVY)

    add_heading(doc, "5.2 Hybrid Candidate Retrieval", 2)
    retrieval_steps = [
        ("BM25 lexical retrieval. ", "Finds exact terms, named entities, abbreviations, and numbers."),
        ("Dense semantic retrieval. ", "Uses Sentence Transformers for meaning-based similarity, with deterministic TF-IDF/SVD as the offline fallback."),
        ("Reciprocal-rank fusion. ", "Combines complementary rankings with RRF constant k = 60."),
        ("Optional cross-encoder. ", "Directly scores question–chunk pairs for stronger top-candidate reranking when the model is available."),
    ]
    number_id = add_numbering_instance(doc, decimal_abstract_id)
    for lead, rest in retrieval_steps:
        add_numbered_item(doc, lead + rest, number_id, bold_lead=lead)

    add_heading(doc, "5.3 Adaptive Feature Engineering", 2)
    add_body(doc, "The controller uses only features available at inference time:")
    for item in [
        "question length, entity and number counts, and comparison, multihop, broad, or calculation indicators;",
        "top relevance score, top-two score gap, normalized score entropy, and largest score-drop position;",
        "Jaccard agreement between BM25 and dense top results; and",
        "query-term coverage, candidate-section dispersion, and available candidate count."
    ]:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "5.4 Learned Context-Budget Controller", 2)
    add_body(
        doc,
        "Gold evidence is mapped to ranked chunks during training, and the smallest depth that recovers "
        "the evidence determines the target budget class. Multinomial logistic regression and a "
        "class-balanced random forest are compared with stratified cross-validation using balanced "
        "accuracy. The better model is serialized with its feature order and training metadata. An "
        "interpretable heuristic keeps the application functional when no model artifact is present."
    )
    add_table(
        doc,
        ["Budget class", "Base chunks", "Typical question pattern"],
        [
            ["Simple", "2", "Direct fact or definition with one concentrated evidence region."],
            ["Moderate", "5", "Explanation, comparison, or answer requiring several nearby passages."],
            ["Complex", "8", "Broad, calculation, multihop, or dispersed-evidence question."],
            ["Expanded maximum", "10", "Low-confidence retrieval or uncertain complex query."],
        ],
        [1900, 1500, 5960],
        first_col_bold=True,
    )

    add_heading(doc, "5.5 Diversity- and Token-Aware Selection", 2)
    add_body(
        doc,
        "A greedy MMR-style selector rewards reranked relevance, novelty, new-section coverage, and "
        "useful neighbouring chunks while penalizing redundancy and token cost. Selection stops when "
        "the assigned budget is filled or additional candidates provide too little marginal value. "
        "Low top-score confidence, or a small score gap combined with high entropy, expands the budget "
        "by one class."
    )
    add_callout(
        doc,
        "Selector utility",
        "U(c|S) = α·relevance + (1−α)·novelty + section bonus + neighbour bonus − token-cost penalty."
    )

    add_heading(doc, "5.6 Answer Generation Modes", 2)
    for text, lead in [
        ("Extractive mode: deterministic, offline, and zero-setup; selects non-redundant evidence sentences.", "Extractive mode:"),
        ("Ollama mode: uses any locally installed chat model through an Ollama server.", "Ollama mode:"),
        ("OpenAI-compatible mode: supports local endpoints such as LM Studio or vLLM and other compatible servers.", "OpenAI-compatible mode:"),
    ]:
        add_list_item(doc, text, bullet_id, bold_lead=lead)
    add_body(
        doc,
        "All modes receive only selected evidence and source identifiers. Grounded prompts prohibit "
        "unsupported knowledge and require chunk citations."
    )

    add_heading(doc, "6. Model Evaluation", 1)
    add_heading(doc, "6.1 Paired Comparison Strategies", 2)
    add_body(
        doc,
        "Every strategy answers the same held-out questions so differences are attributable to context "
        "construction rather than sample composition. The six strategies are full context, BM25 "
        "fixed-k, dense fixed-k, hybrid fixed-k, largest-score-drop elbow, and AdaContext adaptive "
        "selection. The fixed hybrid baseline isolates the value of adaptation from retrieval quality."
    )

    add_heading(doc, "6.2 Metrics", 2)
    for text, lead in [
        ("Answer quality: exact match and maximum token-level F1 over accepted references.", "Answer quality:"),
        ("Evidence quality: gold-evidence recall and context precision.", "Evidence quality:"),
        ("Efficiency: selected chunks, selected context tokens, token reduction, and component/total latency.", "Efficiency:"),
        ("Uncertainty: paired bootstrap 95% confidence intervals over question-level outputs.", "Uncertainty:"),
        ("Robustness: beginning/middle/end evidence-position tests at increasing context lengths.", "Robustness:"),
    ]:
        add_list_item(doc, text, bullet_id, bold_lead=lead)

    add_heading(doc, "6.3 Bundled Demonstration Results", 2)
    add_body(
        doc,
        "The table below reports the included four-question test split. These are demonstration results "
        "that verify the pipeline, not final research evidence."
    )
    add_table(
        doc,
        ["Strategy", "Token F1", "Evidence recall", "Median tokens", "Mean token reduction"],
        [
            ["Full context", "0.2671", "100%", "489.5", "0.0%"],
            ["Hybrid fixed-k", "0.2678", "100%", "296.5", "39.9%"],
            ["Elbow", "0.2603", "100%", "100.5", "70.2%"],
            ["AdaContext", "0.2671", "100%", "301.5", "35.2%"],
        ],
        [2000, 1400, 1700, 1900, 2360],
        first_col_bold=True,
    )
    add_body(
        doc,
        "AdaContext retained all annotated evidence and matched the full-context answer F1 while using "
        "about one-third fewer context tokens. The elbow method was cheaper but slightly lower in F1, "
        "illustrating why quality and cost must be interpreted jointly."
    )

    add_heading(doc, "6.4 Lost-in-the-Middle Stress Test", 2)
    add_table(
        doc,
        ["Context size", "Adaptive recall at start / middle / end", "Adaptive median tokens", "Full-context tokens"],
        [
            ["10 chunks", "100% / 100% / 100%", "278–305", "925"],
            ["25 chunks", "100% / 100% / 100%", "340–350.5", "2,373"],
            ["50 chunks", "83.3% / 83.3% / 83.3%", "262–266", "4,780"],
        ],
        [1550, 3400, 2220, 2190],
        first_col_bold=True,
    )
    add_body(
        doc,
        "Recall was consistent across evidence positions at each context size, indicating that the "
        "retriever did not systematically fail on middle evidence. The drop at 50 chunks reveals a "
        "real scalability limit that the larger held-out experiment should quantify."
    )

    add_heading(doc, "6.5 Final Research Protocol", 2)
    add_body(
        doc,
        "The final report should train on QASPER train records, tune on validation, and evaluate the "
        "chosen configuration on test. Results should be grouped by difficulty, answer type, document "
        "length, and evidence dispersion. A method is preferred when it lies on the quality–cost "
        "Pareto frontier; if two methods differ by no more than two F1 points, the lower-median-token "
        "method is preferred."
    )

    add_heading(doc, "7. Deployment and Monitoring", 1)
    add_heading(doc, "7.1 Streamlit Application", 2)
    add_body(
        doc,
        "The application is a local, CPU-compatible demonstration with four main tabs: Ask, Compare "
        "strategies, Experiments, and Methodology. It can index the bundled sample, QASPER subsets, or "
        "uploaded files; display selected evidence and scores; explain the controller decision; and "
        "compare the answer against baseline context strategies."
    )

    add_heading(doc, "7.2 Custom Questions", 2)
    add_body(
        doc,
        "Yes—custom questions are supported. A question will work when its answer is present or "
        "inferable from the indexed documents. Extractive mode is strongest for facts and short "
        "explanations; a local generative model is better for synthesis, comparison, and calculation. "
        "If the evidence is absent, the grounded system should report insufficient support rather than "
        "inventing an answer."
    )

    add_heading(doc, "7.3 Software, Hardware, and Reproducibility", 2)
    for item in [
        "Python 3.10+ with NumPy, pandas, scikit-learn, PyArrow, PyPDF, Plotly, Streamlit, Sentence Transformers, joblib, and pytest;",
        "normal laptop hardware and CPU-only execution for the bundled workflow;",
        "a versioned YAML configuration containing seeds, chunk sizes, model names, thresholds, baselines, and evaluation settings;",
        "repeatable scripts for controller training, EDA, evaluation, stress testing, and tests; and",
        "question-level outputs that allow aggregate scores to be audited."
    ]:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "7.4 Monitoring", 2)
    add_body(
        doc,
        "Each normal query is appended to a JSONL log containing the question, strategy, predicted "
        "budget, expansion decision, selected chunk identifiers, relevance scores, token counts, and "
        "latency. Monitoring can reveal budget drift, frequent low-confidence expansion, unusually "
        "large prompts, or document types where evidence recall deteriorates."
    )

    add_heading(doc, "7.5 Limitations and Responsible Use", 2)
    limitations = [
        "The default extractive generator is weak at calculations and multi-hop synthesis.",
        "Approximate token counts may differ from a provider-specific tokenizer.",
        "Evidence matching is approximate when annotations cross chunk boundaries.",
        "QASPER focuses on scientific papers and may not generalize to legal, conversational, or code data.",
        "Full-context baselines can exceed a selected model's context window and must be marked when capped.",
        "High-stakes answers require human review because retrieval can still return incomplete or misleading evidence."
    ]
    for item in limitations:
        add_list_item(doc, item, bullet_id)

    add_heading(doc, "Summary", 1)
    add_body(
        doc,
        "AdaContext demonstrates the complete Foundation of Data Science cycle. It begins with a "
        "measurable problem, collects documents and gold evidence, performs structured preprocessing "
        "and EDA, engineers interpretable query/retrieval features, trains a lightweight predictive "
        "controller, evaluates paired baselines with statistical uncertainty, and deploys the result "
        "as a monitored application."
    )
    add_body(
        doc,
        "The project advances beyond a simple cosine-similarity demo by separating retrieval quality "
        "from budget adaptation, using complementary retrieval methods, learning the context size, "
        "penalizing redundancy and token cost, and testing evidence position explicitly. Its expected "
        "outcome is comparable answer quality with materially lower context cost for many questions, "
        "while retaining transparent traces for cases that need more evidence."
    )
    add_callout(
        doc,
        "Submission message",
        "The folder contains the working application, source package, tests, QASPER files, sample data, "
        "notebooks, model artifacts, generated charts, evaluation tables, monitoring logs, and written "
        "methodology. The remaining action is simply to run the workflow and present the results."
    )

    add_heading(doc, "Appendix A — How to Run and Demonstrate", 1)
    add_heading(doc, "A.1 Start the Application", 2)
    add_body(doc, "Open PowerShell in the AdaContext folder and run:")
    for command in [".\\setup.ps1", ".\\run_app.ps1"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(command)
        set_run_font(r, name="Consolas", size=10, color=NAVY, bold=True)

    add_heading(doc, "A.2 Suggested Classroom Walkthrough", 2)
    walkthrough = [
        ("Open Ask. ", "Use a bundled question and show the answer, predicted budget, selected chunks, citations, and token savings."),
        ("Ask a custom question. ", "Choose one whose answer is visibly present in the loaded documents and compare the retrieved evidence."),
        ("Open Compare strategies. ", "Contrast full context, fixed-k, elbow, and adaptive selection for the same question."),
        ("Open Experiments. ", "Show EDA charts, the quality–cost Pareto result, confidence intervals, and the lost-in-the-middle test."),
        ("Open Methodology. ", "Explain the feature set, budget classes, confidence expansion, safeguards, and limitations."),
    ]
    number_id = add_numbering_instance(doc, decimal_abstract_id)
    for lead, rest in walkthrough:
        add_numbered_item(doc, lead + rest, number_id, bold_lead=lead)

    add_heading(doc, "A.3 Reproduce All Local Results", 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(".\\run_all.ps1")
    set_run_font(r, name="Consolas", size=10, color=NAVY, bold=True)
    add_body(
        doc,
        "This retrains the controller, regenerates EDA and evaluation artifacts, runs the controlled "
        "stress test, and executes the automated test suite. Generated outputs are written under "
        "reports/generated/."
    )

    doc.add_page_break()
    add_heading(doc, "Appendix B — Key Project Files", 1)
    add_table(
        doc,
        ["Path", "Purpose"],
        [
            ["app.py", "Interactive Streamlit demonstration and dashboard."],
            ["src/adacontext/", "Chunking, retrieval, features, controller, selection, generation, evaluation, and logging package."],
            ["config/default.yaml", "Reproducible models, thresholds, chunk settings, budgets, and strategies."],
            ["data/sample/", "Bundled four-document, 24-question gold benchmark."],
            ["data/raw/qasper/", "Official QASPER train, validation, and test Parquet files."],
            ["artifacts/", "Serialized budget-controller models and diagnostics."],
            ["reports/generated/", "EDA tables, interactive charts, evaluation summaries, and stress-test outputs."],
            ["logs/query_log.jsonl", "Query-level context decisions, selected evidence, token counts, and latency."],
            ["tests/", "Offline automated tests covering preprocessing, retrieval, pipeline behavior, datasets, and metrics."],
        ],
        [2750, 6610],
        font_size=8.9,
        first_col_bold=True,
        keep_together=False,
    )

    add_heading(doc, "References", 1)
    references = [
        "Dasigi, P., et al. (2021). A Dataset of Information-Seeking Questions and Answers Anchored in Research Papers. NAACL.",
        "Jeong, S., et al. (2024). Adaptive-RAG: Learning to Adapt Retrieval-Augmented Large Language Models through Question Complexity. NAACL.",
        "Bai, Y., et al. (2025). LongBench v2: Towards Deeper Understanding and Reasoning on Realistic Long-context Multitasks. ACL.",
        "Sentence Transformers. Semantic Search documentation. https://sbert.net/examples/sentence_transformer/applications/semantic-search/README.html"
    ]
    number_id = add_numbering_instance(doc, decimal_abstract_id)
    for ref in references:
        add_numbered_item(doc, ref, number_id)

    # Ensure every section follows the same page geometry and quiet running furniture.
    for section in doc.sections:
        configure_page(section)
        configure_header_footer(section)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
